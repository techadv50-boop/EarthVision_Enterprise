"""Author database: under-process records scoped to assigned journals."""

from __future__ import annotations

import pytest
from httpx import AsyncClient


def _bearer(login) -> dict[str, str]:
    return {"Authorization": f"Bearer {login.json()['access_token']}"}


@pytest.mark.asyncio
async def test_admin_adds_under_process_article(client: AsyncClient):
    operator = await client.post(
        "/api/v1/auth/login",
        json={"username": "citation@xdgen.com", "password": "pak123"},
    )
    headers = _bearer(operator)
    created = await client.post(
        "/api/v1/author-articles",
        headers=headers,
        json={
            "ojs_number": "IJIST-2024-118",
            "title": "EdDSA watermarking for documents",
            "author_names": "Ali Khan; Sara Ahmed",
            "author_emails": "ali@example.com; sara@example.com",
            "email_sent_date": "2024-01-05",
            "plagiarism": "9%",
            "orcid_id": "0000-0002-1825-0097",
            "received_date": "2024-01-12",
            "review_rounds": [
                {"round": 1, "sent_date": "2024-01-20", "received_date": "2024-02-01"},
            ],
            "accepted_date": "2024-03-10",
            "galley_sent_date": "2024-03-12",
            "galley_received_date": "2024-03-18",
            "publish_date": "",
            "editorial_status": "Waiting for reviewer to be assigned",
        },
    )
    assert created.status_code == 201, created.text
    body = created.json()
    assert body["wing"] == "in_process"
    assert body["ojs_number"] == "IJIST-2024-118"
    assert body["email_sent_date"] == "2024-01-05"
    assert body["plagiarism"] == "9%"
    assert body["review_rounds"][0]["sent_date"] == "2024-01-20"
    assert body["galley_sent_date"] == "2024-03-12"
    assert body["editorial_status"] == "Waiting for reviewer to be assigned"

    listed = await client.get("/api/v1/author-articles", headers=headers, params={"wing": "in_process"})
    assert listed.status_code == 200
    assert any(row["id"] == body["id"] for row in listed.json())

    published = await client.get("/api/v1/author-articles", headers=headers, params={"wing": "published"})
    assert published.status_code == 200
    assert published.json() == []

    moved = await client.patch(
        f"/api/v1/author-articles/{body['id']}",
        headers=headers,
        json={"wing": "published"},
    )
    assert moved.status_code == 200, moved.text
    assert moved.json()["wing"] == "published"

    still_open = await client.get("/api/v1/author-articles", headers=headers, params={"wing": "in_process"})
    assert all(row["id"] != body["id"] for row in still_open.json())
    now_pub = await client.get("/api/v1/author-articles", headers=headers, params={"wing": "published"})
    assert any(row["id"] == body["id"] for row in now_pub.json())


@pytest.mark.asyncio
async def test_published_article_keeps_published_status(client: AsyncClient):
    operator = await client.post(
        "/api/v1/auth/login",
        json={"username": "citation@xdgen.com", "password": "pak123"},
    )
    headers = _bearer(operator)
    created = await client.post(
        "/api/v1/author-articles",
        headers=headers,
        json={
            "wing": "published",
            "journal_title": "IJIST",
            "ojs_number": "IJIST-PUB-STATUS",
            "title": "Issued paper",
            "author_names": "Hassan Raza",
            "editorial_status": "Published",
            "comments": "Keep this comment box as it is.",
        },
    )
    assert created.status_code == 201, created.text
    assert created.json()["editorial_status"] == "Published"
    assert created.json()["comments"] == "Keep this comment box as it is."


@pytest.mark.asyncio
async def test_assigned_user_only_sees_own_journal_author_records(client: AsyncClient):
    operator = await client.post(
        "/api/v1/auth/login",
        json={"username": "citation@xdgen.com", "password": "pak123"},
    )
    admin = _bearer(operator)
    journal_a = await client.post(
        "/api/v1/journals",
        headers=admin,
        json={"name": "Author DB Alpha", "abbreviation": "ADA"},
    )
    journal_b = await client.post(
        "/api/v1/journals",
        headers=admin,
        json={"name": "Author DB Beta", "abbreviation": "ADB"},
    )
    id_a = journal_a.json()["id"]
    id_b = journal_b.json()["id"]
    created_user = await client.post(
        "/api/v1/admin/users",
        headers=admin,
        json={
            "email": "adbeditor@example.com",
            "username": "adbeditor",
            "password": "EditorPass@123456",
            "role": "user",
            "privileges": {
                "services": ["authors"],
                "review_branches": [],
                "author_wings": ["in_process", "published"],
                "all_journals": False,
            },
            "assigned_journal_ids": [id_a],
        },
    )
    assert created_user.status_code == 201, created_user.text
    login = await client.post(
        "/api/v1/auth/login",
        json={"username": "adbeditor", "password": "EditorPass@123456"},
    )
    user = _bearer(login)

    own = await client.post(
        "/api/v1/author-articles",
        headers=user,
        json={"journal_id": id_a, "ojs_number": "ADA-1", "title": "Alpha paper"},
    )
    assert own.status_code == 201, own.text
    blocked = await client.post(
        "/api/v1/author-articles",
        headers=user,
        json={"journal_id": id_b, "ojs_number": "ADB-1", "title": "Beta paper"},
    )
    assert blocked.status_code == 404

    other = await client.post(
        "/api/v1/author-articles",
        headers=admin,
        json={"journal_id": id_b, "ojs_number": "ADB-9", "title": "Hidden beta"},
    )
    assert other.status_code == 201
    listed = await client.get("/api/v1/author-articles", headers=user, params={"wing": "in_process"})
    ids = {row["id"] for row in listed.json()}
    assert own.json()["id"] in ids
    assert other.json()["id"] not in ids

    deleted = await client.delete(f"/api/v1/author-articles/{own.json()['id']}", headers=user)
    assert deleted.status_code == 403
    still = await client.get("/api/v1/author-articles", headers=user, params={"wing": "in_process"})
    assert any(row["id"] == own.json()["id"] for row in still.json())


@pytest.mark.asyncio
async def test_user_modifications_record_previous_and_new_values(client: AsyncClient):
    operator = await client.post(
        "/api/v1/auth/login",
        json={"username": "citation@xdgen.com", "password": "pak123"},
    )
    headers = _bearer(operator)
    created = await client.post(
        "/api/v1/author-articles",
        headers=headers,
        json={"ojs_number": "MOD-1", "title": "History paper", "received_date": "2024-01-01"},
    )
    assert created.status_code == 201, created.text
    article_id = created.json()["id"]
    first = await client.patch(
        f"/api/v1/author-articles/{article_id}",
        headers=headers,
        json={"received_date": "2024-02-15", "plagiarism": "12%"},
    )
    assert first.status_code == 200, first.text
    mods = first.json()["modifications"]
    assert len(mods) == 1
    assert mods[0]["mod_number"] == 1
    assert mods[0]["account"]
    labels = {item["label"]: item for item in mods[0]["changes"]}
    assert labels["Receive date"]["previous"] == "2024-01-01"
    assert labels["Receive date"]["new"] == "2024-02-15"
    assert labels["Plagiarism"]["previous"] == "(empty)"
    assert labels["Plagiarism"]["new"] == "12%"
    assert first.json()["original_snapshot"]["received_date"] == "2024-01-01"
    assert first.json()["modifications"][0]["snapshot"]["received_date"] == "2024-02-15"
    assert first.json()["modifications"][0]["snapshot"]["plagiarism"] == "12%"

    second = await client.patch(
        f"/api/v1/author-articles/{article_id}",
        headers=headers,
        json={"received_date": "2024-03-20"},
    )
    assert second.status_code == 200
    assert [item["mod_number"] for item in second.json()["modifications"]] == [1, 2]
    latest = second.json()["modifications"][1]
    assert latest["changes"][0]["previous"] == "2024-02-15"
    assert latest["changes"][0]["new"] == "2024-03-20"


@pytest.mark.asyncio
async def test_under_process_review_rounds_and_journal_catalog(client: AsyncClient):
    operator = await client.post(
        "/api/v1/auth/login",
        json={"username": "citation@xdgen.com", "password": "pak123"},
    )
    headers = _bearer(operator)
    journals = await client.get("/api/v1/author-articles/journals", headers=headers)
    assert journals.status_code == 200, journals.text
    names = {row["name"] for row in journals.json()}
    assert "International Journal of Innovations in Science & Technology" in names
    assert "Magna Carta: Contemporary Social Science" in names
    assert "International Journal of Agriculture and Sustainable Development" in names
    assert "Frontiers in Computational Spatial Intelligence" in names
    assert "Journal of International Relations and Social Dynamics" in names
    assert "International Journal of NT Diseases" in names
    assert "Demo Extra Journal" not in names
    assert "Journal A — Innovations in Science & Technology" not in names
    assert "Journal B — Applied Earth Studies" not in names
    assert "Journal C — Hidden from standard users" not in names
    assert "State Update Demo Journal" not in names
    assert "IJIST-A" not in {row.get("abbreviation") for row in journals.json()}
    assert "JAES-B" not in {row.get("abbreviation") for row in journals.json()}
    assert "JHC-C" not in {row.get("abbreviation") for row in journals.json()}
    assert "SUDJ" not in {row.get("abbreviation") for row in journals.json()}
    cite_journals = await client.get("/api/v1/journals", headers=headers)
    assert cite_journals.status_code == 200
    assert "FCSI" not in {row.get("abbreviation") for row in cite_journals.json()}
    created = await client.post(
        "/api/v1/author-articles",
        headers=headers,
        json={
            "journal_title": "IJIST",
            "ojs_number": "RND-1",
            "title": "Round paper",
            "email_sent_date": "2024-06-01",
            "review_rounds": [
                {"sent_date": "2024-06-02", "received_date": "2024-06-10"},
                {"sent_date": "2024-06-12", "received_date": "2024-06-20"},
            ],
            "accepted_date": "2024-07-01",
            "galley_sent_date": "2024-07-02",
            "galley_received_date": "2024-07-08",
        },
    )
    assert created.status_code == 201, created.text
    assert created.json()["journal_id"] is None
    assert created.json()["journal_name"] == "International Journal of Innovations in Science & Technology"
    rounds = created.json()["review_rounds"]
    assert len(rounds) == 2
    assert rounds[1]["round"] == 2
    assert rounds[1]["sent_date"] == "2024-06-12"
    third = await client.patch(
        f"/api/v1/author-articles/{created.json()['id']}",
        headers=headers,
        json={
            "review_rounds": [
                *rounds,
                {"sent_date": "2024-06-22", "received_date": "2024-06-30"},
            ]
        },
    )
    assert third.status_code == 200, third.text
    assert len(third.json()["review_rounds"]) == 3
    assert "Round 3" in third.json()["modifications"][-1]["changes"][0]["new"]


@pytest.mark.asyncio
async def test_excel_import_creates_and_updates_author_articles(client: AsyncClient):
    from io import BytesIO

    from openpyxl import Workbook

    operator = await client.post(
        "/api/v1/auth/login",
        json={"username": "citation@xdgen.com", "password": "pak123"},
    )
    headers = _bearer(operator)
    template = await client.get("/api/v1/author-articles/template", headers=headers)
    assert template.status_code == 200
    assert "spreadsheet" in template.headers.get("content-type", "")

    book = Workbook()
    sheet = book.active
    sheet.append(
        [
            "OJS number",
            "Title",
            "Author names",
            "Email addresses of authors",
            "Email sent",
            "Plagiarism",
            "ORCID ID",
            "Receive date",
            "Review date",
            "Accepted date",
            "Publish date",
            "Repeat done",
            "DOI in PDF",
        ]
    )
    sheet.append(
        [
            "XL-100",
            "Imported paper",
            "Noor Ali",
            "noor@example.com",
            "Yes",
            "8%",
            "0000-0001-2345-6789",
            "2024-04-01",
            "",
            "",
            "",
            "No",
            "10.1/xl-100",
        ]
    )
    buf = BytesIO()
    book.save(buf)
    imported = await client.post(
        "/api/v1/author-articles/import",
        headers=headers,
        params={"wing": "in_process"},
        files={
            "file": (
                "authors.xlsx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert imported.status_code == 200, imported.text
    assert imported.json()["created"] == 1
    article_id = imported.json()["articles"][0]["id"]

    book2 = Workbook()
    sheet2 = book2.active
    sheet2.append(["OJS number", "Title", "Receive date", "Plagiarism"])
    sheet2.append(["XL-100", "Imported paper revised", "2024-05-09", "10%"])
    buf2 = BytesIO()
    book2.save(buf2)
    updated = await client.post(
        "/api/v1/author-articles/import",
        headers=headers,
        params={"wing": "in_process"},
        files={
            "file": (
                "authors2.xlsx",
                buf2.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert updated.status_code == 200, updated.text
    assert updated.json()["updated"] == 1
    detail = (await client.get("/api/v1/author-articles", headers=headers, params={"wing": "in_process"})).json()
    row = next(item for item in detail if item["id"] == article_id)
    assert row["received_date"] == "2024-05-09"
    assert row["plagiarism"] == "10%"
    assert row["modifications"]
    assert row["modifications"][-1]["account"]
    assert row["original_snapshot"]["received_date"] == "2024-04-01"
    assert row["modifications"][-1]["snapshot"]["received_date"] == "2024-05-09"


@pytest.mark.asyncio
async def test_admin_adds_author_journal_name_users_cannot(client: AsyncClient):
    operator = await client.post(
        "/api/v1/auth/login",
        json={"username": "citation@xdgen.com", "password": "pak123"},
    )
    admin = _bearer(operator)
    added = await client.post(
        "/api/v1/author-articles/journals",
        headers=admin,
        json={"name": "New Author Journal of Testing"},
    )
    assert added.status_code == 201, added.text
    listed = await client.get("/api/v1/author-articles/journals", headers=admin)
    names = {row["name"] for row in listed.json()}
    assert "New Author Journal of Testing" in names
    created_user = await client.post(
        "/api/v1/admin/users",
        headers=admin,
        json={
            "email": "adbplus@example.com",
            "username": "adbplus",
            "password": "EditorPass@123456",
            "role": "user",
            "privileges": {
                "services": ["authors"],
                "review_branches": [],
                "author_wings": ["in_process"],
                "all_journals": True,
            },
        },
    )
    assert created_user.status_code == 201, created_user.text
    login = await client.post(
        "/api/v1/auth/login",
        json={"username": "adbplus", "password": "EditorPass@123456"},
    )
    user = _bearer(login)
    blocked = await client.post(
        "/api/v1/author-articles/journals",
        headers=user,
        json={"name": "Secret Journal"},
    )
    assert blocked.status_code == 403


@pytest.mark.asyncio
async def test_admin_deletes_ojs_record_only_with_password(client: AsyncClient):
    operator = await client.post(
        "/api/v1/auth/login",
        json={"username": "citation@xdgen.com", "password": "pak123"},
    )
    headers = _bearer(operator)
    created = await client.post(
        "/api/v1/author-articles",
        headers=headers,
        json={"ojs_number": "DEL-OJS-1", "title": "Delete me", "received_date": "2026-01-01"},
    )
    assert created.status_code == 201, created.text
    article_id = created.json()["id"]
    await client.patch(
        f"/api/v1/author-articles/{article_id}",
        headers=headers,
        json={"plagiarism": "4%"},
    )
    fetched = await client.get(f"/api/v1/author-articles/{article_id}", headers=headers)
    assert fetched.status_code == 200
    assert fetched.json()["ojs_number"] == "DEL-OJS-1"
    assert len(fetched.json()["modifications"]) == 1

    missing = await client.delete(f"/api/v1/author-articles/{article_id}", headers=headers)
    assert missing.status_code == 403
    wrong = await client.request(
        "DELETE",
        f"/api/v1/author-articles/{article_id}",
        headers=headers,
        json={"password": "not-the-password"},
    )
    assert wrong.status_code == 403
    still = await client.get(f"/api/v1/author-articles/{article_id}", headers=headers)
    assert still.status_code == 200

    deleted = await client.request(
        "DELETE",
        f"/api/v1/author-articles/{article_id}",
        headers=headers,
        json={"password": "pak123"},
    )
    assert deleted.status_code == 204, deleted.text
    gone = await client.get(f"/api/v1/author-articles/{article_id}", headers=headers)
    assert gone.status_code == 404
    listed = await client.get("/api/v1/author-articles", headers=headers, params={"wing": "in_process"})
    assert all(row["id"] != article_id for row in listed.json())


@pytest.mark.asyncio
async def test_under_process_reminders_comments_and_current_stage(client: AsyncClient):
    operator = await client.post(
        "/api/v1/auth/login",
        json={"username": "citation@xdgen.com", "password": "pak123"},
    )
    headers = _bearer(operator)
    comment = "Reviewer has not replied. " * 20
    assert len(comment) < 1000
    created = await client.post(
        "/api/v1/author-articles",
        headers=headers,
        json={
            "ojs_number": "SLA-7",
            "title": "Current-state paper",
            "review_rounds": [{"round": 1, "sent_date": "2026-09-01", "received_date": ""}],
            "soft_reminder_sent": "2026-09-05",
            "second_reminder_sent": "2026-09-12",
            "last_reminder_sent": "2026-09-19",
            "comments": comment,
            "current_stage": "round_1_sent",
            "current_stage_started": "2026-09-01",
            "current_stage_days": 7,
            "current_stage_passed": False,
        },
    )
    assert created.status_code == 201, created.text
    body = created.json()
    assert body["soft_reminder_sent"] == "2026-09-05"
    assert body["second_reminder_sent"] == "2026-09-12"
    assert body["last_reminder_sent"] == "2026-09-19"
    assert body["comments"] == comment
    assert body["current_stage"] == "round_1_sent"
    assert body["current_stage_started"] == "2026-09-01"
    assert body["current_stage_days"] == 7
    assert body["current_stage_passed"] is False
    assert body["original_snapshot"]["current_stage"] == "round_1_sent"

    too_long = await client.post(
        "/api/v1/author-articles",
        headers=headers,
        json={"ojs_number": "SLA-LONG", "title": "Too long", "comments": "x" * 1001},
    )
    assert too_long.status_code == 422

    patched = await client.patch(
        f"/api/v1/author-articles/{body['id']}",
        headers=headers,
        json={
            "current_stage_passed": True,
            "comments": "Review returned.",
            "current_stage": "Round 1 review receive date",
            "current_stage_started": "2026-09-20",
        },
    )
    assert patched.status_code == 200, patched.text
    assert patched.json()["current_stage_passed"] is True
    assert patched.json()["current_stage"] == "round_1_received"
    assert patched.json()["comments"] == "Review returned."
    labels = {item["label"]: item for item in patched.json()["modifications"][-1]["changes"]}
    assert labels["Current state"]["previous"] == "Round 1 review sent date"
    assert labels["Current state"]["new"] == "Round 1 review receive date"
    assert labels["Current state passed"]["new"] == "Yes"


@pytest.mark.asyncio
async def test_current_issue_sanitization_blocks_same_authors(client: AsyncClient):
    operator = await client.post(
        "/api/v1/auth/login",
        json={"username": "citation@xdgen.com", "password": "pak123"},
    )
    headers = _bearer(operator)
    journal = "International Journal of Innovations in Science & Technology"
    first = await client.post(
        "/api/v1/author-articles",
        headers=headers,
        json={
            "wing": "published",
            "journal_title": "IJIST",
            "ojs_number": "IJIST-ISSUE-1",
            "title": "Already in the issue",
            "author_names": "Ali Khan; Noor Ali",
            "author_emails": "ali@example.com",
        },
    )
    assert first.status_code == 201, first.text
    other = await client.post(
        "/api/v1/author-articles",
        headers=headers,
        json={
            "wing": "published",
            "journal_title": "IJIST",
            "ojs_number": "IJIST-ISSUE-2",
            "title": "Second paper in the issue",
            "author_names": "Sara Ahmed",
        },
    )
    assert other.status_code == 201, other.text
    saved = await client.put(
        "/api/v1/author-articles/sanitization",
        headers=headers,
        json={
            "journal_title": "IJIST",
            "label": "Current issue",
            "article_ids": [first.json()["id"], other.json()["id"]],
        },
    )
    assert saved.status_code == 200, saved.text
    assert saved.json()["journal_title"] == journal
    assert set(saved.json()["article_ids"]) == {first.json()["id"], other.json()["id"]}
    assert "Ali Khan" in saved.json()["authors"]
    assert "Sara Ahmed" in saved.json()["authors"]
    assert all(item["wing"] == "published" for item in saved.json()["published"])
    assert {first.json()["id"], other.json()["id"]}.issubset(
        {item["id"] for item in saved.json()["published"]}
    )

    scheduled = await client.post(
        "/api/v1/author-articles",
        headers=headers,
        json={
            "journal_title": "IJIST",
            "ojs_number": "IJIST-NEXT",
            "title": "Scheduled for the same issue",
            "author_names": "Khan, Ali",
        },
    )
    assert scheduled.status_code == 201, scheduled.text
    listing = await client.get(
        "/api/v1/author-articles/sanitization",
        headers=headers,
        params={"journal_title": "IJIST"},
    )
    assert listing.status_code == 200
    assert scheduled.json()["id"] in {item["id"] for item in listing.json()["scheduled"]}
    assert scheduled.json()["id"] not in {item["id"] for item in listing.json()["published"]}
    assert all(item["wing"] == "in_process" for item in listing.json()["scheduled"])
    rejected = await client.post(
        "/api/v1/author-articles/sanitization/check",
        headers=headers,
        json={"article_id": first.json()["id"]},
    )
    assert rejected.status_code == 400
    checked = await client.post(
        "/api/v1/author-articles/sanitization/check",
        headers=headers,
        json={
            "article_id": scheduled.json()["id"],
            "article_ids": [first.json()["id"], other.json()["id"]],
        },
    )
    assert checked.status_code == 200, checked.text
    assert checked.json()["allowed"] is False
    assert checked.json()["overlaps"]
    blocked = await client.patch(
        f"/api/v1/author-articles/{scheduled.json()['id']}",
        headers=headers,
        json={"wing": "published"},
    )
    assert blocked.status_code == 409, blocked.text
    assert "same issue" in blocked.json()["detail"].lower()
    publish = await client.post(
        "/api/v1/author-articles/sanitization/publish",
        headers=headers,
        json={"article_id": scheduled.json()["id"]},
    )
    assert publish.status_code == 409

    clean = await client.post(
        "/api/v1/author-articles",
        headers=headers,
        json={
            "journal_title": "IJIST",
            "ojs_number": "IJIST-CLEAN",
            "title": "Different authors",
            "author_names": "Hassan Raza",
        },
    )
    assert clean.status_code == 201, clean.text
    ok = await client.post(
        "/api/v1/author-articles/sanitization/check",
        headers=headers,
        json={"article_id": clean.json()["id"]},
    )
    assert ok.status_code == 200
    assert ok.json()["allowed"] is True
    published = await client.post(
        "/api/v1/author-articles/sanitization/publish",
        headers=headers,
        json={"article_id": clean.json()["id"]},
    )
    assert published.status_code == 200, published.text
    assert published.json()["allowed"] is True
    basket = await client.get(
        "/api/v1/author-articles/sanitization",
        headers=headers,
        params={"journal_title": "IJIST"},
    )
    assert clean.json()["id"] in basket.json()["article_ids"]


@pytest.mark.asyncio
async def test_excel_export_includes_published_and_under_process_rows(client: AsyncClient):
    from io import BytesIO

    from openpyxl import load_workbook

    operator = await client.post(
        "/api/v1/auth/login",
        json={"username": "citation@xdgen.com", "password": "pak123"},
    )
    headers = _bearer(operator)
    created = await client.post(
        "/api/v1/author-articles",
        headers=headers,
        json={
            "wing": "in_process",
            "journal_title": "IJIST",
            "ojs_number": "EXP-IN-1",
            "title": "Export under process paper",
            "author_names": "Noor Ali",
            "author_emails": "noor@example.com",
            "received_date": "2024-04-01",
        },
    )
    assert created.status_code == 201, created.text
    published = await client.post(
        "/api/v1/author-articles",
        headers=headers,
        json={
            "wing": "published",
            "journal_title": "IJIST",
            "ojs_number": "EXP-PUB-1",
            "title": "Export published paper",
            "author_names": "Export Only Author",
            "author_emails": "",
            "comments": "Article URL: https://example.test/article/view/88",
            "editorial_status": "Published",
            "publish_date": "2024-08-01",
        },
    )
    assert published.status_code == 201, published.text

    under = await client.get(
        "/api/v1/author-articles/export",
        headers=headers,
        params={"wing": "in_process"},
    )
    assert under.status_code == 200, under.text
    assert "spreadsheet" in under.headers.get("content-type", "")
    book = load_workbook(BytesIO(under.content))
    sheet = book.active
    headers_row = [cell for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    assert "OJS number" in headers_row
    assert "Article URL" in headers_row
    assert "Volume" in headers_row
    assert "Issue" in headers_row
    assert "Page" in headers_row
    assert "Reason of decline" in headers_row
    ojs_idx = headers_row.index("OJS number")
    values = [row[ojs_idx] for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert "EXP-IN-1" in values

    pub = await client.get(
        "/api/v1/author-articles/export",
        headers=headers,
        params={"wing": "published"},
    )
    assert pub.status_code == 200, pub.text
    pub_book = load_workbook(BytesIO(pub.content))
    pub_sheet = pub_book.active
    pub_headers = [cell for cell in next(pub_sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    url_idx = pub_headers.index("Article URL")
    ojs_idx = pub_headers.index("OJS number")
    email_idx = pub_headers.index("Email addresses of authors")
    rows = list(pub_sheet.iter_rows(min_row=2, values_only=True))
    found = next(row for row in rows if row[ojs_idx] == "EXP-PUB-1")
    assert found[url_idx] == "https://example.test/article/view/88"
    assert not (found[email_idx] or "").strip()


@pytest.mark.asyncio
async def test_published_archive_crawl_fills_article_fields(client: AsyncClient, monkeypatch):
    import io
    import re

    from reportlab.pdfgen import canvas

    from tests.test_citation_parser import GALLEY_WATER
    from app.services import crawler as crawler_mod
    from app.services.author_crawler import run_author_crawl_job, to_iso_date

    assert to_iso_date("July 22, 2026") == "2026-07-22"
    assert to_iso_date("2026-08-23") == "2026-08-23"

    def _pdf_from_text(text: str) -> bytes:
        buf = io.BytesIO()
        c = canvas.Canvas(buf)
        y = 800
        for line in text.strip().splitlines():
            c.drawString(40, y, line[:110])
            y -= 14
            if y < 40:
                c.showPage()
                y = 800
        c.save()
        return buf.getvalue()

    pdf_bytes = _pdf_from_text(GALLEY_WATER)
    operator = await client.post(
        "/api/v1/auth/login",
        json={"username": "citation@xdgen.com", "password": "pak123"},
    )
    headers = _bearer(operator)

    async def fake_fetch(url: str):
        pages = {
            "https://example.test/issue/archive": (
                b"<html><a href='https://example.test/issue/view/1'>Vol. 8 No. 5 (2026)</a></html>"
            ),
            "https://example.test/issue/view/1": (
                b"<html>"
                b'<a href="https://example.test/article/view/9">Water</a>'
                b'<a href="https://example.test/article/view/9/11">PDF</a>'
                b'<a href="https://example.test/article/view/10">Second</a>'
                b"</html>"
            ),
            "https://example.test/article/view/9": (
                b"<html>"
                b'<meta name="citation_title" content="Integrated Source Tracking and Assessment of Drinking Water Contamination in Rural Sindh, Pakistan">'
                b'<meta name="citation_author" content="Asim Ali">'
                b'<meta name="citation_author" content="Jabir Ali Keerio">'
                b'<meta name="citation_doi" content="10.33411/IJIST/202608052211">'
                b'<meta name="citation_pdf_url" content="https://example.test/article/download/9/11">'
                b"IJIST-2026-2211"
                b"</html>"
            ),
            "https://example.test/article/view/10": (
                b"<html>"
                b'<meta name="citation_title" content="A second published paper from the archive">'
                b'<meta name="citation_author" content="Hassan Raza">'
                b'<meta name="citation_publication_date" content="2026-08-01">'
                b"IJIST-2026-4010"
                b"</html>"
            ),
        }
        if url in pages:
            return 200, pages[url], "text/html"
        if url.endswith(".pdf") or "/download/" in url or re.search(r"/article/view/\d+/\d+", url):
            return 200, pdf_bytes, "application/pdf"
        return 404, b"", "text/plain"

    monkeypatch.setattr(crawler_mod, "default_fetch", fake_fetch)

    start = await client.post(
        "/api/v1/author-articles/crawl",
        headers=headers,
        json={
            "journal_title": "IJIST",
            "archive_url": "https://example.test/issue/archive",
        },
    )
    assert start.status_code == 200, start.text
    job_id = start.json()["id"]
    job = (await client.get(f"/api/v1/author-articles/crawl/{job_id}", headers=headers)).json()
    if job["status"] in {"queued", "running", "failed"} or int(job.get("articles_saved") or 0) < 2:
        await run_author_crawl_job(job_id, fetch=fake_fetch)
        job = (await client.get(f"/api/v1/author-articles/crawl/{job_id}", headers=headers)).json()
    assert job["status"] == "completed", job
    listed = await client.get(
        "/api/v1/author-articles",
        headers=headers,
        params={"wing": "published"},
    )
    assert listed.status_code == 200
    rows = listed.json()
    water = next(row for row in rows if "Drinking Water" in (row.get("title") or "") or row.get("ojs_number") == "IJIST-2026-2211")
    assert water["wing"] == "published"
    assert water["editorial_status"] == "Published"
    assert water["volume"] == 8
    assert water["issue"] == 5
    assert str(water["page"]).startswith("2211")
    assert "Asim" in (water["author_names"] or "") or "Ali" in (water["author_names"] or "")
    assert water["publish_date"]
    assert "example.test/article/view/9" in (water.get("comments") or "")
    second = next(row for row in rows if row.get("ojs_number") == "IJIST-2026-4010" or "second published paper" in (row.get("title") or "").lower())
    assert second["title"]
    assert second["author_names"]

    again = await run_author_crawl_job(job_id, fetch=fake_fetch)
    del again
    dup = (await client.get(f"/api/v1/author-articles/crawl/{job_id}", headers=headers)).json()
    assert dup["articles_already"] >= 2 or dup["articles_saved"] == 0

    missing = await client.get("/api/v1/webcrawler/status", headers=headers)
    assert missing.status_code == 404


@pytest.mark.asyncio
async def test_users_cannot_move_articles_to_published(client: AsyncClient):
    operator = await client.post(
        "/api/v1/auth/login",
        json={"username": "citation@xdgen.com", "password": "pak123"},
    )
    admin = _bearer(operator)
    created_user = await client.post(
        "/api/v1/admin/users",
        headers=admin,
        json={
            "email": "voluser@example.com",
            "username": "voluser",
            "password": "EditorPass@123456",
            "role": "user",
            "privileges": {
                "services": ["authors"],
                "review_branches": [],
                "author_wings": ["in_process", "published"],
                "all_journals": True,
            },
        },
    )
    assert created_user.status_code == 201, created_user.text
    login = await client.post(
        "/api/v1/auth/login",
        json={"username": "voluser", "password": "EditorPass@123456"},
    )
    user = _bearer(login)
    created = await client.post(
        "/api/v1/author-articles",
        headers=user,
        json={
            "ojs_number": "VOL-MOVE-1",
            "title": "User cannot publish directly",
            "author_names": "Unique Vol Author",
            "journal_title": "IJIST",
        },
    )
    assert created.status_code == 201, created.text
    blocked = await client.patch(
        f"/api/v1/author-articles/{created.json()['id']}",
        headers=user,
        json={"wing": "published"},
    )
    assert blocked.status_code == 403, blocked.text
    assert created.json()["id"] == (await client.get(
        f"/api/v1/author-articles/{created.json()['id']}", headers=user
    )).json()["id"]
    still = await client.get(
        f"/api/v1/author-articles/{created.json()['id']}",
        headers=user,
    )
    assert still.json()["wing"] == "in_process"

    moved = await client.patch(
        f"/api/v1/author-articles/{created.json()['id']}",
        headers=admin,
        json={"wing": "published"},
    )
    assert moved.status_code == 200, moved.text
    assert moved.json()["wing"] == "published"

    scheduled = await client.post(
        "/api/v1/author-articles",
        headers=user,
        json={
            "ojs_number": "VOL-SAN-1",
            "title": "Sanitization can still publish",
            "author_names": "Sanitization Only Author",
            "journal_title": "IJIST",
        },
    )
    assert scheduled.status_code == 201, scheduled.text
    published = await client.post(
        "/api/v1/author-articles/sanitization/publish",
        headers=user,
        json={"article_id": scheduled.json()["id"]},
    )
    assert published.status_code == 200, published.text
    assert published.json()["scheduled"]["wing"] == "published"


@pytest.mark.asyncio
async def test_published_article_stores_volume_issue_page(client: AsyncClient):
    operator = await client.post(
        "/api/v1/auth/login",
        json={"username": "citation@xdgen.com", "password": "pak123"},
    )
    headers = _bearer(operator)
    created = await client.post(
        "/api/v1/author-articles",
        headers=headers,
        json={
            "wing": "published",
            "journal_title": "IJIST",
            "ojs_number": "IJIST-VOL-8-5",
            "title": "Volume grouped paper",
            "author_names": "Volume Author",
            "editorial_status": "Published",
            "volume": 8,
            "issue": 5,
            "page": "1788-1813",
        },
    )
    assert created.status_code == 201, created.text
    body = created.json()
    assert body["volume"] == 8
    assert body["issue"] == 5
    assert body["page"] == "1788-1813"
    listed = await client.get("/api/v1/author-articles", headers=headers, params={"wing": "published"})
    match = next(row for row in listed.json() if row["id"] == body["id"])
    assert match["volume"] == 8
    assert match["issue"] == 5
    assert match["page"] == "1788-1813"


@pytest.mark.asyncio
async def test_declined_status_stores_reason(client: AsyncClient):
    operator = await client.post(
        "/api/v1/auth/login",
        json={"username": "citation@xdgen.com", "password": "pak123"},
    )
    headers = _bearer(operator)
    created = await client.post(
        "/api/v1/author-articles",
        headers=headers,
        json={
            "ojs_number": "IJIST-DECLINED-1",
            "title": "Declined paper",
            "author_names": "Declined Author",
            "editorial_status": "Declined",
            "decline_reason": "Outside the journal scope.",
        },
    )
    assert created.status_code == 201, created.text
    assert created.json()["editorial_status"] == "Declined"
    assert created.json()["decline_reason"] == "Outside the journal scope."
    patched = await client.patch(
        f"/api/v1/author-articles/{created.json()['id']}",
        headers=headers,
        json={"editorial_status": "rejected", "decline_reason": "Methods are not reproducible."},
    )
    assert patched.status_code == 200, patched.text
    assert patched.json()["editorial_status"] == "Declined"
    assert patched.json()["decline_reason"] == "Methods are not reproducible."


@pytest.mark.asyncio
async def test_excel_template_matches_new_form_fields(client: AsyncClient):
    from io import BytesIO

    from openpyxl import load_workbook

    operator = await client.post(
        "/api/v1/auth/login",
        json={"username": "citation@xdgen.com", "password": "pak123"},
    )
    headers = _bearer(operator)
    template = await client.get("/api/v1/author-articles/template", headers=headers)
    assert template.status_code == 200
    book = load_workbook(BytesIO(template.content))
    sheet = book.active
    headers_row = [cell for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    assert headers_row.index("OJS number") == 0
    assert headers_row[1:4] == ["Volume", "Issue", "Page"]
    assert "Status" in headers_row
    assert headers_row[headers_row.index("Status") + 1] == "Reason of decline"
    sample = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))
    assert len(sample) >= len(headers_row) or len([cell for cell in sample if cell is not None]) > 0
    assert sample[1] == "8" or sample[1] == 8
    assert sample[2] == "5" or sample[2] == 5
    assert str(sample[3]) == "1788-1813"


