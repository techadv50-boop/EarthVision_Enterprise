"""SatPass Predict helpers: parse KML/shapefile uploads and export pass reports."""

from __future__ import annotations

import io
import json
import tempfile
import zipfile
from pathlib import Path
from typing import Annotated, Any, Literal

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

from app.core.dependencies import get_current_user
from app.models.user import User

router = APIRouter(prefix="/satellites/predict", tags=["SatPass Predict"])


class ExportRequest(BaseModel):
    format: Literal["csv", "xlsx", "pdf"]
    title: str = "SatPass prediction"
    headers: list[str] = Field(default_factory=list)
    rows: list[list[str]] = Field(default_factory=list)


def _features_from_geojson(data: dict[str, Any]) -> dict[str, Any]:
    if data.get("type") == "FeatureCollection":
        return data
    if data.get("type") == "Feature":
        return {"type": "FeatureCollection", "features": [data]}
    return {
        "type": "FeatureCollection",
        "features": [{"type": "Feature", "properties": {}, "geometry": data}],
    }


def _parse_kml_simple(raw: bytes) -> dict[str, Any]:
    """Minimal KML Point/Polygon/LineString reader used when Fiona has no KML driver."""
    import xml.etree.ElementTree as ET

    def local(tag: str) -> str:
        return tag.split("}")[-1].lower()

    def coords(text: str | None) -> list[list[float]]:
        out: list[list[float]] = []
        for token in (text or "").replace("\n", " ").split():
            parts = token.split(",")
            if len(parts) >= 2:
                try:
                    out.append([float(parts[0]), float(parts[1])])
                except ValueError:
                    continue
        return out

    root = ET.fromstring(raw)
    features: list[dict[str, Any]] = []
    for el in root.iter():
        if local(el.tag) != "placemark":
            continue
        name = ""
        geom: dict[str, Any] | None = None
        for child in el.iter():
            kind = local(child.tag)
            if kind == "name" and child.text and not name:
                name = child.text.strip()
            elif kind == "point":
                for c in child.iter():
                    if local(c.tag) == "coordinates":
                        pts = coords(c.text)
                        if pts:
                            geom = {"type": "Point", "coordinates": pts[0]}
            elif kind == "linestring":
                for c in child.iter():
                    if local(c.tag) == "coordinates":
                        pts = coords(c.text)
                        if len(pts) >= 2:
                            geom = {"type": "LineString", "coordinates": pts}
            elif kind == "polygon":
                ring: list[list[float]] = []
                for c in child.iter():
                    if local(c.tag) == "coordinates":
                        ring = coords(c.text)
                        break
                if len(ring) >= 3:
                    if ring[0] != ring[-1]:
                        ring.append(ring[0])
                    geom = {"type": "Polygon", "coordinates": [ring]}
        if geom:
            features.append({"type": "Feature", "properties": {"name": name}, "geometry": geom})
    if not features:
        raise ValueError("No placemarks with coordinates found in KML.")
    return {"type": "FeatureCollection", "features": features}


def _parse_kml_bytes(raw: bytes) -> dict[str, Any]:
    try:
        import geopandas as gpd

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "upload.kml"
            path.write_bytes(raw)
            gdf = gpd.read_file(path)
            if gdf.crs and str(gdf.crs).upper() not in {"EPSG:4326", "WGS84"}:
                gdf = gdf.to_crs(4326)
            parsed = json.loads(gdf.to_json())
            if parsed.get("features"):
                return parsed
    except Exception:
        pass
    try:
        return _parse_kml_simple(raw)
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail="Could not parse KML. Try a GeoJSON/shapefile, or a simpler KML.",
        ) from exc


@router.post("/geometry")
async def parse_geometry(
    _user: Annotated[User, Depends(get_current_user)],
    files: list[UploadFile] = File(...),
):
    """Parse uploaded KML or a complete shapefile set (.shp/.shx/.dbf/.prj or .zip) to GeoJSON."""
    if not files:
        raise HTTPException(status_code=400, detail="No files uploaded.")

    names = [(f.filename or "upload").lower() for f in files]
    blobs = [await f.read() for f in files]

    if len(files) == 1 and (names[0].endswith(".kml") or names[0].endswith(".kmz")):
        if names[0].endswith(".kmz"):
            with zipfile.ZipFile(io.BytesIO(blobs[0])) as zf:
                kml_name = next((n for n in zf.namelist() if n.lower().endswith(".kml")), None)
                if not kml_name:
                    raise HTTPException(status_code=400, detail="KMZ did not contain a KML file.")
                raw = zf.read(kml_name)
        else:
            raw = blobs[0]
        geojson = _parse_kml_bytes(raw)
        geojson = _features_from_geojson(geojson)
        return {"geojson": geojson, "feature_count": len(geojson.get("features", []))}

    if len(files) == 1 and (names[0].endswith(".geojson") or names[0].endswith(".json")):
        try:
            data = json.loads(blobs[0].decode("utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise HTTPException(status_code=400, detail="Invalid GeoJSON.") from exc
        geojson = _features_from_geojson(data)
        return {"geojson": geojson, "feature_count": len(geojson.get("features", []))}

    try:
        import geopandas as gpd
    except ImportError as exc:
        raise HTTPException(status_code=500, detail="GeoPandas is not installed.") from exc

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        if len(files) == 1 and names[0].endswith(".zip"):
            zip_path = tmp_path / "upload.zip"
            zip_path.write_bytes(blobs[0])
            gdf = gpd.read_file(zip_path)
        else:
            for name, blob in zip(names, blobs):
                (tmp_path / Path(name).name).write_bytes(blob)
            shp = next(tmp_path.glob("*.shp"), None)
            if shp is None:
                raise HTTPException(
                    status_code=400,
                    detail="Upload a .zip shapefile or the full set (.shp, .shx, .dbf, optional .prj).",
                )
            gdf = gpd.read_file(shp)
        if gdf.empty:
            raise HTTPException(status_code=400, detail="The shapefile contained no features.")
        if gdf.crs and str(gdf.crs).upper() not in {"EPSG:4326", "WGS84"}:
            gdf = gdf.to_crs(4326)
        geojson = json.loads(gdf.to_json())
        return {"geojson": geojson, "feature_count": len(gdf)}


@router.post("/export")
async def export_report(
    payload: ExportRequest,
    _user: Annotated[User, Depends(get_current_user)],
):
    fmt = payload.format
    title = payload.title.strip() or "SatPass prediction"
    headers = payload.headers
    rows = payload.rows
    safe = "".join(ch if ch.isalnum() or ch in "-_." else "_" for ch in title)[:80] or "satpass-prediction"
    if fmt == "csv":
        buf = io.StringIO()
        import csv

        writer = csv.writer(buf)
        writer.writerow([title])
        writer.writerow([])
        writer.writerow(headers)
        writer.writerows(rows)
        data = buf.getvalue().encode("utf-8")
        return StreamingResponse(
            io.BytesIO(data),
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="{safe}.csv"'},
        )

    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "Passes"
        ws["A1"] = title
        ws["A1"].font = Font(size=14, bold=True)
        fill = PatternFill(start_color="0E7490", end_color="0E7490", fill_type="solid")
        for col, h in enumerate(headers, start=1):
            cell = ws.cell(row=3, column=col, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = fill
        for r_i, row in enumerate(rows, start=4):
            for c_i, val in enumerate(row, start=1):
                ws.cell(row=r_i, column=c_i, value=val)
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return StreamingResponse(
            out,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{safe}.xlsx"'},
        )

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=18, rightMargin=18)
    styles = getSampleStyleSheet()
    elements = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
    table_data = [headers or ["Column"]] + rows
    table = Table(table_data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0e7490")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    elements.append(table)
    doc.build(elements)
    out.seek(0)
    return StreamingResponse(
        out,
        media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="{safe}.pdf"'},
    )
