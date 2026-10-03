"""Author database management: under-process and published articles."""

from __future__ import annotations

import io
import re
from datetime import datetime, timezone
from typing import Annotated, Any, Optional

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core.dependencies import get_current_user
from app.database.session import get_db
from app.models.citation import AuthorArticle, AuthorArticleChange, Journal
from app.models.user import User
from app.schemas.author_db import (
    FIELD_LABELS,
    AuthorArticleIn,
    AuthorArticleOut,
    AuthorArticlePatch,
    AuthorFieldChangeOut,
    AuthorImportResult,
    AuthorModificationOut,
    WINGS,
)
from app.services.journal_access import allowed_journal_ids, require_journal_access

router = APIRouter(prefix="/author-articles", tags=["Author database"])

Db = Annotated[AsyncSession, Depends(get_db)]
CurrentUser = Annotated[User, Depends(get_current_user)]

EXCEL_HEADERS = [
    "OJS number",
    "Title",
    "Author names",
    "Email addresses of authors",
    "Email sent",
    "Plagiarism",
    "ORCID ID",
    "Receive date",
    "Review date",
    "Accepted date",
    "Publish date",
    "Repeat done",
    "DOI in PDF",
    "Journal",
    "Wing",
]


def _blank(value: Optional[str]) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _display(value: Any) -> str:
    if value is None:
        return "(empty)"
    if isinstance(value, bool):
        return "Yes" if value else "No"
    text = str(value).strip()
    return text or "(empty)"


def _account_label(user: User) -> str:
    name = (user.full_name or "").strip()
    email = (user.email or "").strip()
    username = (user.username or "").strip()
    if name and email:
        return f"{name} ({email})"
    if email:
        return email
    return username or f"user #{user.id}"


def _modification_out(change: AuthorArticleChange) -> AuthorModificationOut:
    raw = change.changes or []
    fields = [
        AuthorFieldChangeOut(
            field=str(item.get("field") or ""),
            label=str(item.get("label") or item.get("field") or ""),
            previous=str(item.get("previous") if item.get("previous") is not None else "(empty)"),
            new=str(item.get("new") if item.get("new") is not None else "(empty)"),
        )
        for item in raw
        if isinstance(item, dict)
    ]
    return AuthorModificationOut(
        mod_number=change.mod_number,
        changed_at=change.changed_at,
        account=change.account_name or change.account_email or change.account_username or "Unknown account",
        account_username=change.account_username or "",
        account_email=change.account_email or "",
        account_name=change.account_name or "",
        changes=fields,
    )


def _payload(row: AuthorArticle) -> AuthorArticleOut:
    journal = row.journal
    mods = [_modification_out(item) for item in sorted(row.changes or [], key=lambda item: item.mod_number)]
    return AuthorArticleOut(
        id=row.id,
        wing=row.wing,
        journal_id=row.journal_id,
        journal_name=(journal.abbreviation or journal.name) if journal else None,
        owner_id=row.owner_id,
        ojs_number=row.ojs_number or "",
        title=row.title or "",
        author_names=row.author_names or "",
        author_emails=row.author_emails or "",
        email_sent=bool(row.email_sent),
        plagiarism=row.plagiarism or "",
        orcid_id=row.orcid_id or "",
        received_date=row.received_date,
        review_date=row.review_date,
        accepted_date=row.accepted_date,
        publish_date=row.publish_date,
        repeat_done=bool(row.repeat_done),
        doi_in_pdf=row.doi_in_pdf or "",
        created_at=row.created_at,
        updated_at=row.updated_at,
        modifications=mods,
    )


def _load_options():
    return (selectinload(AuthorArticle.journal), selectinload(AuthorArticle.changes))


async def _visible_query(db: AsyncSession, user: User, wing: Optional[str]):
    stmt = select(AuthorArticle).options(*_load_options()).order_by(AuthorArticle.id.desc())
    if wing:
        stmt = stmt.where(AuthorArticle.wing == wing)
    allowed = await allowed_journal_ids(db, user)
    if allowed is not None:
        if not allowed:
            stmt = stmt.where(AuthorArticle.owner_id == user.id, AuthorArticle.journal_id.is_(None))
        else:
            stmt = stmt.where(
                or_(AuthorArticle.journal_id.in_(allowed), AuthorArticle.owner_id == user.id)
            )
    return stmt


async def _load_row(db: AsyncSession, article_id: int, user: User) -> AuthorArticle:
    row = (
        await db.execute(
            select(AuthorArticle).options(*_load_options()).where(AuthorArticle.id == article_id)
        )
    ).scalar_one_or_none()
    if row is None:
        raise HTTPException(status_code=404, detail="Article not found")
    allowed = await allowed_journal_ids(db, user)
    if allowed is None:
        return row
    if row.owner_id == user.id:
        return row
    if row.journal_id and row.journal_id in allowed:
        return row
    raise HTTPException(status_code=404, detail="Article not found")


async def _journal_label(db: AsyncSession, journal_id: Optional[int]) -> str:
    if not journal_id:
        return "(empty)"
    journal = await db.get(Journal, journal_id)
    if journal is None:
        return str(journal_id)
    return journal.abbreviation or journal.name or str(journal_id)


async def _record_modifications(
    db: AsyncSession,
    row: AuthorArticle,
    user: User,
    diffs: list[dict[str, str]],
) -> None:
    if not diffs:
        return
    nxt = int(
        (
            await db.execute(
                select(func.max(AuthorArticleChange.mod_number)).where(
                    AuthorArticleChange.article_id == row.id
                )
            )
        ).scalar()
        or 0
    ) + 1
    db.add(
        AuthorArticleChange(
            article_id=row.id,
            mod_number=nxt,
            changed_at=datetime.now(timezone.utc),
            user_id=user.id,
            account_username=user.username or "",
            account_email=user.email or "",
            account_name=_account_label(user),
            changes=diffs,
        )
    )


async def _apply_update(
    db: AsyncSession,
    row: AuthorArticle,
    user: User,
    data: dict[str, Any],
) -> list[dict[str, str]]:
    diffs: list[dict[str, str]] = []
    if "wing" in data and data["wing"] is not None:
        wing = str(data["wing"]).strip().lower()
        if wing not in WINGS:
            raise HTTPException(status_code=400, detail="Wing must be in_process or published")
        if row.wing != wing:
            diffs.append(
                {
                    "field": "wing",
                    "label": FIELD_LABELS["wing"],
                    "previous": "Under process" if row.wing == "in_process" else "Published",
                    "new": "Under process" if wing == "in_process" else "Published",
                }
            )
            row.wing = wing
    if "journal_id" in data:
        journal_id = data["journal_id"]
        if journal_id is not None:
            await require_journal_access(db, user, int(journal_id))
            journal_id = int(journal_id)
        if row.journal_id != journal_id:
            diffs.append(
                {
                    "field": "journal_id",
                    "label": FIELD_LABELS["journal_id"],
                    "previous": await _journal_label(db, row.journal_id),
                    "new": await _journal_label(db, journal_id),
                }
            )
            row.journal_id = journal_id
    for field in (
        "ojs_number",
        "title",
        "author_names",
        "author_emails",
        "plagiarism",
        "orcid_id",
        "doi_in_pdf",
    ):
        if field in data and data[field] is not None:
            nxt = str(data[field]).strip()
            prev = getattr(row, field) or ""
            if prev != nxt:
                diffs.append(
                    {
                        "field": field,
                        "label": FIELD_LABELS[field],
                        "previous": _display(prev),
                        "new": _display(nxt),
                    }
                )
                setattr(row, field, nxt)
    for field in ("received_date", "review_date", "accepted_date", "publish_date"):
        if field in data:
            nxt = _blank(data[field] if data[field] is None else str(data[field]))
            prev = getattr(row, field)
            if (prev or None) != nxt:
                diffs.append(
                    {
                        "field": field,
                        "label": FIELD_LABELS[field],
                        "previous": _display(prev),
                        "new": _display(nxt),
                    }
                )
                setattr(row, field, nxt)
    for field in ("email_sent", "repeat_done"):
        if field in data and data[field] is not None:
            nxt = bool(data[field])
            prev = bool(getattr(row, field))
            if prev != nxt:
                diffs.append(
                    {
                        "field": field,
                        "label": FIELD_LABELS[field],
                        "previous": _display(prev),
                        "new": _display(nxt),
                    }
                )
                setattr(row, field, nxt)
    if diffs:
        row.updated_at = datetime.now(timezone.utc)
        await _record_modifications(db, row, user, diffs)
    return diffs


async def _reload(db: AsyncSession, article_id: int) -> AuthorArticle:
    return (
        await db.execute(
            select(AuthorArticle).options(*_load_options()).where(AuthorArticle.id == article_id)
        )
    ).scalar_one()


@router.get("", response_model=list[AuthorArticleOut])
async def list_author_articles(
    db: Db,
    user: CurrentUser,
    wing: Optional[str] = Query(default=None),
):
    if wing and wing not in WINGS:
        raise HTTPException(status_code=400, detail="Wing must be in_process or published")
    rows = list((await db.execute(await _visible_query(db, user, wing))).scalars().all())
    return [_payload(row) for row in rows]


@router.get("/template")
async def download_import_template():
    book = Workbook()
    sheet = book.active
    sheet.title = "Author database"
    sheet.append(EXCEL_HEADERS)
    sheet.append(
        [
            "IJIST-2024-118",
            "Sample title",
            "Ali Khan; Sara Ahmed",
            "ali@example.com; sara@example.com",
            "Yes",
            "9%",
            "0000-0002-1825-0097",
            "2024-01-12",
            "2024-02-01",
            "2024-03-10",
            "",
            "No",
            "10.33411/IJIST/20240101118",
            "IJIST",
            "in_process",
        ]
    )
    buf = io.BytesIO()
    book.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="author-database-template.xlsx"'},
    )


@router.post("", response_model=AuthorArticleOut, status_code=status.HTTP_201_CREATED)
async def create_author_article(body: AuthorArticleIn, db: Db, user: CurrentUser):
    wing = (body.wing or "in_process").strip().lower()
    if wing not in WINGS:
        raise HTTPException(status_code=400, detail="Wing must be in_process or published")
    if body.journal_id is not None:
        await require_journal_access(db, user, body.journal_id)
    row = AuthorArticle(
        wing=wing,
        journal_id=body.journal_id,
        owner_id=user.id,
        ojs_number=(body.ojs_number or "").strip(),
        title=(body.title or "").strip(),
        author_names=(body.author_names or "").strip(),
        author_emails=(body.author_emails or "").strip(),
        email_sent=bool(body.email_sent),
        plagiarism=(body.plagiarism or "").strip(),
        orcid_id=(body.orcid_id or "").strip(),
        received_date=_blank(body.received_date),
        review_date=_blank(body.review_date),
        accepted_date=_blank(body.accepted_date),
        publish_date=_blank(body.publish_date),
        repeat_done=bool(body.repeat_done),
        doi_in_pdf=(body.doi_in_pdf or "").strip(),
    )
    if not row.title and not row.ojs_number:
        raise HTTPException(status_code=400, detail="Enter an OJS number or a title.")
    db.add(row)
    await db.flush()
    return _payload(await _reload(db, row.id))


@router.patch("/{article_id}", response_model=AuthorArticleOut)
async def update_author_article(
    article_id: int, body: AuthorArticlePatch, db: Db, user: CurrentUser
):
    row = await _load_row(db, article_id, user)
    await _apply_update(db, row, user, body.model_dump(exclude_unset=True))
    await db.flush()
    db.expire(row, ["changes", "journal"])
    return _payload(await _reload(db, row.id))


@router.delete("/{article_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_author_article(article_id: int, db: Db, user: CurrentUser):
    if not user.is_citation_admin():
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Users cannot delete author-database records. Save a modification instead.",
        )
    row = await _load_row(db, article_id, user)
    await db.delete(row)
    await db.flush()
    return None


def _norm_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").strip().lower()).strip()


HEADER_MAP = {
    "ojs": "ojs_number",
    "ojs number": "ojs_number",
    "ojs no": "ojs_number",
    "title": "title",
    "author names": "author_names",
    "authors": "author_names",
    "author name": "author_names",
    "email": "author_emails",
    "emails": "author_emails",
    "email addresses": "author_emails",
    "email addresses of authors": "author_emails",
    "author emails": "author_emails",
    "email sent": "email_sent",
    "if the email sent": "email_sent",
    "plagiarism": "plagiarism",
    "orcid": "orcid_id",
    "orcid id": "orcid_id",
    "receive date": "received_date",
    "received date": "received_date",
    "review date": "review_date",
    "accepted date": "accepted_date",
    "accept date": "accepted_date",
    "publish date": "publish_date",
    "published date": "publish_date",
    "repeat done": "repeat_done",
    "repeat": "repeat_done",
    "doi": "doi_in_pdf",
    "doi in pdf": "doi_in_pdf",
    "journal": "journal",
    "wing": "wing",
}


def _parse_bool(value: Any) -> bool:
    text = str(value or "").strip().lower()
    return text in {"1", "y", "yes", "true", "done", "sent"}


def _parse_date(value: Any) -> Optional[str]:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    text = str(value).strip()
    if not text:
        return None
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        return text
    if isinstance(value, float):
        try:
            from openpyxl.utils.datetime import from_excel

            return from_excel(value).date().isoformat()
        except Exception:
            return text
    return text[:32]


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    return str(value).strip()


def _read_table(data: bytes, filename: str) -> list[dict[str, str]]:
    name = (filename or "").lower()
    if name.endswith(".csv"):
        import csv

        text = data.decode("utf-8-sig", errors="replace")
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
    else:
        book = load_workbook(io.BytesIO(data), data_only=True)
        sheet = book.active
        rows = [[cell for cell in line] for line in sheet.iter_rows(values_only=True)]
    if not rows:
        return []
    headers = [_norm_header(col) for col in rows[0]]
    mapped = [HEADER_MAP.get(header) for header in headers]
    if not any(mapped):
        raise HTTPException(
            status_code=400,
            detail="That file does not use the author-database column names. Download the Excel template.",
        )
    out: list[dict[str, str]] = []
    for raw in rows[1:]:
        if raw is None:
            continue
        item: dict[str, str] = {}
        empty = True
        for idx, key in enumerate(mapped):
            if not key or idx >= len(raw):
                continue
            value = raw[idx]
            if key in ("received_date", "review_date", "accepted_date", "publish_date"):
                parsed = _parse_date(value)
                item[key] = parsed or ""
            elif key in ("email_sent", "repeat_done"):
                item[key] = "yes" if _parse_bool(value) else "no"
            else:
                item[key] = _cell(value)
            if item[key]:
                empty = False
        if not empty:
            out.append(item)
    return out


async def _journal_by_name(db: AsyncSession, user: User, name: str) -> Optional[int]:
    needle = name.strip().lower()
    if not needle:
        return None
    journals = list((await db.execute(select(Journal))).scalars().all())
    allowed = await allowed_journal_ids(db, user)
    for journal in journals:
        label = (journal.abbreviation or journal.name or "").strip().lower()
        full = (journal.name or "").strip().lower()
        if needle in {label, full} and (allowed is None or journal.id in allowed):
            return journal.id
    return None


@router.post("/import", response_model=AuthorImportResult)
async def import_author_articles(
    db: Db,
    user: CurrentUser,
    file: UploadFile = File(...),
    wing: str = Query(default="in_process"),
):
    wing = (wing or "in_process").strip().lower()
    if wing not in WINGS:
        raise HTTPException(status_code=400, detail="Wing must be in_process or published")
    data = await file.read()
    if not data:
        raise HTTPException(status_code=400, detail="The Excel file is empty.")
    try:
        table = _read_table(data, file.filename or "import.xlsx")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Could not read that Excel file. {exc}") from exc

    created = 0
    updated = 0
    skipped = 0
    errors: list[str] = []
    articles: list[AuthorArticleOut] = []
    for index, item in enumerate(table, start=2):
        title = item.get("title") or ""
        ojs = item.get("ojs_number") or ""
        if not title and not ojs:
            skipped += 1
            continue
        journal_id = None
        if item.get("journal"):
            journal_id = await _journal_by_name(db, user, item["journal"])
            if journal_id is None:
                errors.append(f"Row {index}: journal “{item['journal']}” was not found or is not assigned.")
                skipped += 1
                continue
        wanted_wing = (item.get("wing") or wing).strip().lower()
        if wanted_wing in {"under process", "under-process", "in process"}:
            wanted_wing = "in_process"
        if wanted_wing in {"published articles", "published"}:
            wanted_wing = "published"
        if wanted_wing not in WINGS:
            wanted_wing = wing
        existing = None
        if ojs:
            stmt = select(AuthorArticle).options(*_load_options()).where(
                AuthorArticle.ojs_number == ojs, AuthorArticle.wing == wanted_wing
            )
            if journal_id is not None:
                stmt = stmt.where(AuthorArticle.journal_id == journal_id)
            existing = (await db.execute(stmt.order_by(AuthorArticle.id.desc()))).scalars().first()
            if existing is not None:
                try:
                    await _load_row(db, existing.id, user)
                except HTTPException:
                    existing = None
        payload = {
            "wing": wanted_wing,
            "journal_id": journal_id if journal_id is not None else (existing.journal_id if existing else None),
            "ojs_number": ojs,
            "title": title,
            "author_names": item.get("author_names") or "",
            "author_emails": item.get("author_emails") or "",
            "email_sent": _parse_bool(item.get("email_sent")),
            "plagiarism": item.get("plagiarism") or "",
            "orcid_id": item.get("orcid_id") or "",
            "received_date": item.get("received_date") or None,
            "review_date": item.get("review_date") or None,
            "accepted_date": item.get("accepted_date") or None,
            "publish_date": item.get("publish_date") or None,
            "repeat_done": _parse_bool(item.get("repeat_done")),
            "doi_in_pdf": item.get("doi_in_pdf") or "",
        }
        try:
            if existing is None:
                created_row = await create_author_article(AuthorArticleIn(**payload), db, user)
                created += 1
                articles.append(created_row)
            else:
                await _apply_update(db, existing, user, payload)
                await db.flush()
                db.expire(existing, ["changes", "journal"])
                updated += 1
                articles.append(_payload(await _reload(db, existing.id)))
        except HTTPException as exc:
            errors.append(f"Row {index}: {exc.detail}")
            skipped += 1
    return AuthorImportResult(
        created=created,
        updated=updated,
        skipped=skipped,
        errors=errors,
        articles=articles,
    )
